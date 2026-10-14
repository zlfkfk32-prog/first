"""Extract only TOP100 rank/name fields from the official CAK civil/building sheet."""
import argparse
import csv
from pathlib import Path
import re
import sys

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from cardletter.rankings import DEFAULT_PATH, load_roster


def main():
    parser = argparse.ArgumentParser(description="대한건설협회 토건 시트의 공개 1~100위 명단 추출")
    parser.add_argument("--xlsx", type=Path, required=True)
    parser.add_argument("--year", type=int, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--previous", type=Path, default=DEFAULT_PATH, help="명칭이 같은 회사의 기존 별칭 유지")
    args = parser.parse_args()
    workbook = load_workbook(args.xlsx, read_only=True, data_only=True)
    try:
        sheet = workbook["토건"]
        title = str(sheet["A1"].value)
        if str(args.year) not in title or "토목건축공사업" not in title:
            parser.error("원본 토건 시트의 발표 연도/공종을 확인하세요")
        rows = [row for row in sheet.iter_rows(values_only=True) if type(row[0]) is int and 1 <= row[0] <= 100]
        if len(rows) != 100 or {row[0] for row in rows} != set(range(1, 101)):
            parser.error("공식 원본에 1~100위가 빠짐없이 한 번씩 있어야 합니다")
    finally:
        workbook.close()
    previous = {}
    if args.previous.exists():
        try:
            previous = {c.name: c.aliases for c in load_roster(args.previous).contractors}
        except ValueError:
            pass  # A pending roster contains no usable aliases.
    result = []
    for row in rows:
        name = re.sub(r"\s+", "", re.sub(r"(?:주식회사|\(주\)|㈜)", "", row[1]))
        result.append((row[0], name, "|".join(previous.get(name, ()))))
    args.output.parent.mkdir(parents=True, exist_ok=True)
    if args.output.exists():
        parser.error("기존 출력 파일은 덮어쓰지 않습니다. 새 경로를 지정하세요")
    with args.output.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(["rank", "name", "aliases"])
        writer.writerows(result)
    print(f"Extracted: {args.output} (100 ranks/names; no representative/contact fields)")


if __name__ == "__main__":
    main()
